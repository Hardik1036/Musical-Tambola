"""
Module 2.1: Asset Manager
Handles song catalog metadata import (.xlsx, .csv, .docx, .txt),
fuzzy audio file mapping against assets/songs/, and local database persistence.
"""
import os
import re
import csv
import difflib
import logging
from pathlib import Path
from typing import List, Dict, Tuple, Optional, Any

import openpyxl
from docx import Document

from config import SONGS_DIR, SUPPORTED_AUDIO_EXTENSIONS
import database

logger = logging.getLogger("tambola.asset_manager")


class AssetManager:
    """Manages audio track catalog and file pairing."""

    def __init__(self, audio_dir: Optional[Path] = None):
        self.audio_dir = Path(audio_dir) if audio_dir else SONGS_DIR
        self.audio_dir.mkdir(parents=True, exist_ok=True)

    def scan_audio_files(self) -> Dict[str, Path]:
        """Scan target folder for all supported audio files.
        Returns a dict mapping filename lower (e.g., '1.mp3', 'chaiyya chaiyya.mp3') to Path.
        """
        files = {}
        if not self.audio_dir.exists():
            return files
        
        for p in self.audio_dir.iterdir():
            if p.is_file() and p.suffix.lower() in SUPPORTED_AUDIO_EXTENSIONS:
                files[p.name.lower()] = p
                files[p.stem.lower()] = p
        return files

    def parse_file(self, file_path: str) -> List[Dict[str, Any]]:
        """Universal import parser for .xlsx, .csv, .docx, and .txt."""
        path = Path(file_path)
        ext = path.suffix.lower()

        if ext == ".csv":
            return self._parse_csv(path)
        elif ext in [".xlsx", ".xlsm"]:
            return self._parse_excel(path)
        elif ext == ".docx":
            return self._parse_docx(path)
        elif ext in [".txt", ".tsv"]:
            return self._parse_txt(path)
        else:
            raise ValueError(f"Unsupported file format '{ext}'. Allowed: .xlsx, .csv, .docx, .txt")

    def _parse_csv(self, path: Path) -> List[Dict[str, Any]]:
        records = []
        with open(path, "r", encoding="utf-8-sig", errors="replace") as f:
            sample = f.read(2048)
            f.seek(0)
            delimiter = ","
            try:
                sniffer = csv.Sniffer()
                dialect = sniffer.sniff(sample)
                delimiter = dialect.delimiter
            except Exception:
                pass
            
            reader = csv.reader(f, delimiter=delimiter)
            rows = [r for r in reader if any(field.strip() for field in r)]
            if not rows:
                return []

            # Check if first row is header
            first_row = [c.strip().lower() for c in rows[0]]
            id_idx, title_idx, artist_idx, year_idx = -1, -1, -1, -1
            
            for idx, col in enumerate(first_row):
                if any(k in col for k in ["id", "track", "no", "number", "sr", "serial"]):
                    id_idx = idx
                elif any(k in col for k in ["title", "song", "name", "track name"]):
                    title_idx = idx
                elif any(k in col for k in ["artist", "singer", "album", "movie"]):
                    artist_idx = idx
                elif any(k in col for k in ["year", "release"]):
                    year_idx = idx

            start_row = 1 if (id_idx != -1 or title_idx != -1) else 0
            if id_idx == -1 and title_idx == -1:
                # Default assume col 0 is ID, col 1 is Title
                id_idx, title_idx = 0, 1

            for row_num, row in enumerate(rows[start_row:], start=1):
                if not row or not any(row):
                    continue
                song_id = self._extract_id(row[id_idx] if id_idx < len(row) else None, default=row_num)
                title = row[title_idx].strip() if title_idx < len(row) else f"Song #{song_id}"
                artist = row[artist_idx].strip() if 0 <= artist_idx < len(row) else ""
                year = row[year_idx].strip() if 0 <= year_idx < len(row) else ""
                if 1 <= song_id <= 90 and title:
                    records.append({"id": song_id, "title": title, "artist": artist, "year": year})

        return self._normalize_records(records)

    def _parse_excel(self, path: Path) -> List[Dict[str, Any]]:
        wb = openpyxl.load_workbook(path, data_only=True)
        sheet = wb.active
        records = []
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []

        # Find header
        header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
        id_idx, title_idx, artist_idx, year_idx = -1, -1, -1, -1
        for idx, col in enumerate(header):
            if any(k in col for k in ["id", "track", "no", "number", "sr", "serial"]):
                id_idx = idx
            elif any(k in col for k in ["title", "song", "name"]):
                title_idx = idx
            elif any(k in col for k in ["artist", "album", "movie", "singer"]):
                artist_idx = idx
            elif any(k in col for k in ["year"]):
                year_idx = idx

        start_row = 1 if (id_idx != -1 or title_idx != -1) else 0
        if id_idx == -1 and title_idx == -1:
            id_idx, title_idx = 0, 1

        for row_num, row in enumerate(rows[start_row:], start=1):
            if not row or all(v is None for v in row):
                continue
            id_val = row[id_idx] if id_idx < len(row) else None
            song_id = self._extract_id(id_val, default=row_num)
            title = str(row[title_idx]).strip() if title_idx < len(row) and row[title_idx] is not None else f"Song #{song_id}"
            artist = str(row[artist_idx]).strip() if 0 <= artist_idx < len(row) and row[artist_idx] is not None else ""
            year = str(row[year_idx]).strip() if 0 <= year_idx < len(row) and row[year_idx] is not None else ""
            if 1 <= song_id <= 90 and title:
                records.append({"id": song_id, "title": title, "artist": artist, "year": year})

        return self._normalize_records(records)

    def _parse_docx(self, path: Path) -> List[Dict[str, Any]]:
        doc = Document(path)
        records = []

        # 1. Try tables
        for table in doc.tables:
            for row_idx, row in enumerate(table.rows):
                cells = [c.text.strip() for c in row.cells]
                if not cells or not any(cells):
                    continue
                if row_idx == 0 and any("title" in c.lower() or "song" in c.lower() for c in cells):
                    continue  # skip header
                
                if len(cells) >= 2:
                    song_id = self._extract_id(cells[0], default=len(records) + 1)
                    title = cells[1]
                    artist = cells[2] if len(cells) > 2 else ""
                    if 1 <= song_id <= 90 and title:
                        records.append({"id": song_id, "title": title, "artist": artist, "year": ""})
                elif len(cells) == 1:
                    parsed = self._parse_line_text(cells[0], default_id=len(records) + 1)
                    if parsed:
                        records.append(parsed)

        # 2. If no table records, parse paragraphs
        if not records:
            for para in doc.paragraphs:
                text = para.text.strip()
                if not text:
                    continue
                parsed = self._parse_line_text(text, default_id=len(records) + 1)
                if parsed:
                    records.append(parsed)

        return self._normalize_records(records)

    def _parse_txt(self, path: Path) -> List[Dict[str, Any]]:
        records = []
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            for line_idx, line in enumerate(f, start=1):
                text = line.strip()
                if not text or text.startswith("#"):
                    continue
                parsed = self._parse_line_text(text, default_id=len(records) + 1)
                if parsed:
                    records.append(parsed)
        return self._normalize_records(records)

    def _parse_line_text(self, text: str, default_id: int) -> Optional[Dict[str, Any]]:
        # Match patterns like:
        # "1. Kal Ho Naa Ho - Sonu Nigam"
        # "05 - Chaiyya Chaiyya"
        # "12: Jai Ho"
        # "17 | Badtameez Dil"
        match = re.match(r"^(\d{1,3})[\.\-\:\)\|\t\s]+(.*)$", text)
        if match:
            song_id = int(match.group(1))
            rem = match.group(2).strip()
        else:
            song_id = default_id
            rem = text

        if not rem:
            return None

        # Check for artist split like "Song - Artist" or "Song / Movie"
        parts = re.split(r"[\-\|\/]", rem, maxsplit=1)
        title = parts[0].strip()
        artist = parts[1].strip() if len(parts) > 1 else ""

        if 1 <= song_id <= 90 and title:
            return {"id": song_id, "title": title, "artist": artist, "year": ""}
        return None

    def _extract_id(self, val: Any, default: int) -> int:
        if val is None:
            return default
        if isinstance(val, (int, float)):
            return int(val)
        digits = re.findall(r"\d+", str(val))
        if digits:
            return int(digits[0])
        return default

    def _normalize_records(self, records: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        """Deduplicate records by ID and sort 1..90."""
        seen = {}
        for r in records:
            sid = r["id"]
            if 1 <= sid <= 90:
                seen[sid] = r
        return [seen[i] for i in sorted(seen.keys())]

    def map_audio_files(self, catalog: Optional[List[Dict[str, Any]]] = None) -> List[Dict[str, Any]]:
        """Pair catalog items with files in assets/songs/.
        1. Exact numeric index match (1.mp3, 01.wav, track_1.mp3, etc.)
        2. Fuzzy title match
        """
        if catalog is None:
            catalog = database.get_all_songs()

        available_files = self.scan_audio_files()
        file_keys = list(available_files.keys())
        updated = []

        for item in catalog:
            song_id = item["id"]
            title = item.get("title", "").strip().lower()
            matched_path = None

            # 1. Exact numeric matching
            id_patterns = [
                f"{song_id}",
                f"{song_id:02d}",
                f"{song_id:03d}",
                f"track{song_id}",
                f"track_{song_id}",
                f"song{song_id}",
                f"song_{song_id}",
                f"audio{song_id}",
            ]
            
            for pat in id_patterns:
                for ext in SUPPORTED_AUDIO_EXTENSIONS:
                    candidate = pat + ext
                    if candidate in available_files:
                        matched_path = str(available_files[candidate])
                        break
                if matched_path:
                    break

            # 2. Fuzzy title matching if not matched by ID
            if not matched_path and title and file_keys:
                # Compare title against file stems
                clean_title = re.sub(r"[^\w\s]", "", title)
                matches = difflib.get_close_matches(clean_title, file_keys, n=1, cutoff=0.6)
                if matches:
                    matched_path = str(available_files[matches[0]])

            # Keep existing path if already valid and none newly found
            existing = item.get("audio_path", "")
            if not matched_path and existing and Path(existing).exists():
                matched_path = existing

            item["audio_path"] = matched_path or ""
            # Save to database
            database.update_audio_path(song_id, item["audio_path"])
            updated.append(item)

        return updated

    def import_and_save(self, file_path: str, auto_map: bool = True) -> Tuple[int, int]:
        """Parse file, upsert to SQLite, and optionally map audio.
        Returns: (total_imported, mapped_count)
        """
        parsed = self.parse_file(file_path)
        for item in parsed:
            database.upsert_song(
                song_id=item["id"],
                title=item["title"],
                artist=item.get("artist", ""),
                year=item.get("year", ""),
            )

        if auto_map:
            updated = self.map_audio_files()
            mapped = sum(1 for s in updated if s.get("audio_path"))
        try:
            import song_bank
            song_bank.sync_with_database()
        except Exception:
            pass

        return len(parsed), mapped
